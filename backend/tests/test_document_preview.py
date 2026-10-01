from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import Employee, ManagedDocument, Role
from app.services.google_drive import google_drive_worklog_service
from app.services.spreadsheet_preview import preview_spreadsheet


def workbook_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "11509-53"
    sheet.append(["日期", "工時", "金額", "試算"])
    sheet.append([date(2026, 9, 7), 24, 18000, "=B2*750"])
    sheet["E2"] = 0.05
    sheet["E2"].number_format = "0.00%"
    sheet["A3"] = "<script>alert(1)</script>"
    sheet["T45"] = "右下資料"
    workbook.create_sheet("參數")["A1"] = "費率"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_preview_paged_values_formula_warning_and_no_file_mutation():
    content = workbook_bytes()
    data = preview_spreadsheet(content, "計價.xlsm", None, 1, 1)
    assert data["sheets"] == ["11509-53", "參數"]
    assert data["total_rows"] == 45 and data["total_columns"] == 20
    assert len(data["rows"]) == 40 and len(data["columns"]) == 16
    assert data["rows"][1]["cells"][0]["text"] == "2026-09-07T00:00:00".replace("T", " ")
    assert data["rows"][1]["cells"][2]["text"] == "18,000"
    assert data["rows"][1]["cells"][3]["uncached_formula"]
    assert data["rows"][1]["cells"][4]["text"] == "5.00%"
    assert data["has_uncached_formulas"]
    second = preview_spreadsheet(content, "計價.xlsm", "11509-53", 41, 17)
    assert second["columns"] == ["Q", "R", "S", "T"]
    assert second["rows"][-1]["cells"][-1]["text"] == "右下資料"
    original = openpyxl.load_workbook(BytesIO(content))
    assert original["11509-53"]["D2"].value == "=B2*750"
    original.close()


@pytest.mark.parametrize("file_name,sheet,row,column", [
    ("old.xls", None, 1, 1), ("test.xlsx", "不存在", 1, 1),
    ("test.xlsx", None, 100, 1), ("test.xlsx", None, 1, 30),
])
def test_preview_rejects_unsupported_ranges(file_name, sheet, row, column):
    with pytest.raises(ValueError):
        preview_spreadsheet(workbook_bytes(), file_name, sheet, row, column)


def test_invalid_workbook_is_friendly_error():
    with pytest.raises(ValueError, match="格式"):
        preview_spreadsheet(b"not an excel file", "test.xlsx", None, 1, 1)


class TestPreviewAPI:
    def setup_method(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.admin = Employee(employee_code="ADMIN001", name="管理員", role=Role.admin, bind_token="test-admin")
        self.worker = Employee(employee_code="EMP001", name="操作員", role=Role.employee, bind_token="test-worker")
        self.document = ManagedDocument(category="推高機計價", title="計價", original_file_name="計價.xlsm",
                                        stored_file_name="計價.xlsm", drive_file_id="existing-id",
                                        drive_folder_id="existing-folder", drive_url="https://drive.google.com/file/d/existing-id/view")
        self.session.add_all([self.admin, self.worker, self.document])
        self.session.commit()
        app.dependency_overrides[get_session] = lambda: self.session
        app.dependency_overrides[get_current_actor] = lambda: self.admin
        self.client = TestClient(app)

    def teardown_method(self):
        app.dependency_overrides.clear()
        self.session.close()
        self.engine.dispose()

    def test_preview_is_read_only_and_restricted(self):
        with patch.object(google_drive_worklog_service, "download_file_bytes", new=AsyncMock(return_value=workbook_bytes())) as download:
            response = self.client.get(f"/api/documents/{self.document.id}/preview")
            assert response.status_code == 200
            download.assert_awaited_once_with("existing-id")
            assert len(self.session.exec(select(ManagedDocument)).all()) == 1
            app.dependency_overrides[get_current_actor] = lambda: self.worker
            assert self.client.get(f"/api/documents/{self.document.id}/preview").status_code == 403
            assert download.await_count == 1

    def test_errors_and_query_limits(self):
        assert self.client.get("/api/documents/999/preview").status_code == 404
        assert self.client.get(f"/api/documents/{self.document.id}/preview?row_start=0").status_code == 422
        with patch.object(google_drive_worklog_service, "download_file_bytes", new=AsyncMock(side_effect=RuntimeError("private error"))):
            response = self.client.get(f"/api/documents/{self.document.id}/preview")
            assert response.status_code == 502
            assert "private error" not in response.text
