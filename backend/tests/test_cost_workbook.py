from datetime import date
from io import BytesIO

import openpyxl

from app.services.cost_workbook import import_cost_hours, list_cost_sheet_targets


def _workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "11509-53"
    worksheet["A5"] = date(2026, 9, 1)
    worksheet["A6"] = date(2026, 9, 2)
    worksheet["B6"] = 8
    worksheet["C6"] = "=B6*1000"
    worksheet["D6"] = 1
    worksheet["F6"] = 2
    other = workbook.create_sheet("11509-47")
    other["A6"] = date(2026, 9, 2)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_cost_targets_find_roc_month_sheet_and_date_row():
    targets = list_cost_sheet_targets(_workbook_bytes(), "計價表.xlsx", date(2026, 9, 2))

    assert [(target.label, target.worksheet_name, target.row_number) for target in targets] == [
        ("47", "11509-47", 6),
        ("53", "11509-53", 6),
    ]
    target = next(item for item in targets if item.label == "53")
    assert (target.normal_hours, target.overtime_hours, target.support_hours) == (8, 1, 2)


def test_cost_import_only_updates_hour_input_columns_and_keeps_formula():
    content, target = import_cost_hours(
        _workbook_bytes(),
        "計價表.xlsx",
        date(2026, 9, 2),
        "53",
        normal_hours=16,
        overtime_hours=4,
        support_hours=0,
    )

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    worksheet = workbook["11509-53"]
    assert target.row_number == 6
    assert worksheet["B6"].value == 16
    assert worksheet["D6"].value == 4
    assert worksheet["F6"].value == 0
    assert worksheet["C6"].value == "=B6*1000"
