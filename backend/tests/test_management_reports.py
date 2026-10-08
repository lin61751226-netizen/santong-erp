from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import (
    AdminAuditLog, AttendanceEvent, BusinessContact, DocumentDataRevision, Employee,
    FinanceEntry, LeaveRequest, LeaveStatus, ManagedDocument, ManagementWorkbookSnapshot,
    Role, WorkAssignment, Worksite,
)
from app.routes import management_reports as routes
from app.services.management_reports import KINDS, analyze_workbook, employee_matches, report
from app.services.finance_import import finance_dedupe_key


def workbook_bytes():
    w = openpyxl.Workbook()
    s = w.active
    s.title = '115年09月_收支明細表'
    s.append(['收支'])
    for row, values in enumerate([
        [datetime(2026, 9, 1), '', '', '', '上月結餘', '', '', 100, None],
        [datetime(2026, 9, 2), '轉帳', '', '', '銀行', '', '', None, 40],
        [datetime(2026, 9, 3), '付款', '', '未分類', '廠商', '', '', None, 20],
        [datetime(2026, 9, 4), '收款', '工程款', '53', '客戶', '現金', 'AB1', 6000, None],
    ], 4):
        for col, value in enumerate(values, 1): s.cell(row, col, value)
    s = w.create_sheet('115年09月_薪資表')
    for col, value in enumerate([datetime(2026, 9, 1), 1, '林育弘', '堆高機', '司機', 55000, 3000, 2000, 2600, None, None, None, None, None, None, 5000], 1):
        s.cell(6, col, value)
    s['Q6'], s['R6'], s['S6'] = '=SUM(F6:K6)', '=SUM(L6:P6)', '=Q6-R6'
    s['C7'] = '薪轉文字說明'
    s['I7'] = '無薪日'
    s = w.create_sheet('115年09月_排休表')
    s['B6'], s['E6'], s['F6'], s['H6'], s['AI6'] = '林育弘', '加1', '病假', '例休', 'INVALID31'
    s = w.create_sheet('115年09月_行事曆記錄')
    s['F6'] = '9月4日月曆工作'
    s['L8'], s['H8'] = '會議', datetime(2026, 9, 10)
    s['L9'] = '缺日期'
    s = w.create_sheet('公司通訊錄')
    s['B5'], s['C5'], s['G5'], s['J5'] = '測試公司', '12345678', '0912000000', '客戶'
    stream = BytesIO()
    w.save(stream)
    return stream.getvalue()


@pytest.fixture
def ctx():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)
    session = Session(engine)
    actor = Employee(employee_code='ADMIN001', name='測試管理員', bind_token='admin', role=Role.admin,
                     password_hash='unchanged', line_user_id='line-preserved')
    employee = Employee(employee_code='EMP001', name='林育弘', bind_token='driver', role=Role.employee)
    doc = ManagedDocument(category='通訊錄與年度管理', title='全年', original_file_name='annual.xlsx',
                          stored_file_name='annual.xlsx', drive_file_id='source', drive_folder_id='folder', drive_url='https://drive.example/source')
    site = Worksite(code='53', name='齊裕53')
    session.add_all([actor, employee, doc, site])
    session.commit()
    data = analyze_workbook(workbook_bytes())
    snapshot = ManagementWorkbookSnapshot(version_key='initial', source_document_id=doc.id,
        content_sha256=data['content_sha256'], employee_mapping={'林育弘': employee.employee_code}, data=data, created_by_id=actor.id)
    session.add(snapshot)
    session.commit()
    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[get_current_actor] = lambda: actor
    download = AsyncMock(return_value=workbook_bytes())
    backup = AsyncMock(return_value={'status': 'saved'})
    with patch.object(routes.drive, 'download_file_bytes', download), patch.object(routes.drive, 'backup_database', backup):
        yield SimpleNamespace(session=session, actor=actor, employee=employee, doc=doc, site=site,
            data=data, snapshot=snapshot, client=TestClient(app), backup=backup)
    app.dependency_overrides.clear()
    session.close()
    engine.dispose()


def test_parser_payroll_formula_roster_calendar_and_evidence():
    data = analyze_workbook(workbook_bytes())
    assert len(data['payroll']) == 1
    assert data['payroll'][0]['gross'] == 62600
    assert data['payroll'][0]['deductions'] == 5000
    assert data['payroll'][0]['net'] == 57600
    assert data['source_notes'][0]['name'] == '薪轉文字說明'
    assert '31' not in data['roster'][0]['marks']
    assert data['roster'][0]['marks']['1'] == '加1'
    assert {x['date'] for x in data['calendar']} == {'2026-09-04', '2026-09-10'}
    assert [x['status'] for x in data['finance'][:3]] == ['結餘（非營收）', '轉帳（非營收）', '待確認']
    assert len(data['contacts']) == 1


def test_formula_inputs_never_become_zero():
    w = openpyxl.load_workbook(BytesIO(workbook_bytes()))
    w['115年09月_薪資表']['I6'] = '=UNKNOWN()'
    b = BytesIO(); w.save(b)
    data = analyze_workbook(b.getvalue())
    assert data['payroll'][0]['net'] is None
    assert any('公式' in row['reason'] for row in data['warnings'])


@pytest.mark.parametrize('kind', list(KINDS))
def test_all_seven_reports_and_column_alignment(ctx, kind):
    response = ctx.client.get('/api/management-reports/report', params={'month': '2026-09', 'kind': kind, 'snapshot_id': ctx.snapshot.id})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data['rows']
    assert all(len(row) == len(data['fields']) for row in data['rows'])
    assert ctx.actor.password_hash == 'unchanged' and ctx.actor.line_user_id == 'line-preserved'


def test_report_without_snapshot_does_not_invent_salary(ctx):
    data = report(ctx.session, None, '2026-09', 'payroll')
    assert data['rows'] == []
    assert report(ctx.session, None, '2026-09', 'annual')['rows'][8][6] is None


def test_month_year_totals_separate_cash_and_system(ctx):
    ctx.session.add_all([
        FinanceEntry(entry_date=date(2026, 9, 4), entry_type='收入', category='工程款', amount=6000),
        FinanceEntry(entry_date=date(2025, 9, 4), entry_type='收入', category='工程款', amount=999999),
        FinanceEntry(entry_date=date(2026, 10, 1), entry_type='支出', category='薪資', amount=2000),
    ])
    ctx.session.commit()
    data = report(ctx.session, ctx.snapshot, '2026-09', 'annual')
    row = data['rows'][8]
    assert row[1:7] == [6000, 0, 6000, 6100, 60, 57600]
    assert data['summary']['年度系統收入'] == 6000
    assert data['summary']['年度系統支出'] == 2000


def test_finance_manual_edit_alias_not_double_counted(ctx):
    source = ctx.data['finance'][3]
    item = FinanceEntry(entry_date=date(2026, 9, 4), entry_type='收入', category='工程款', amount=7000,
                        income_amount=7000, dedupe_key='new-key')
    ctx.session.add(item); ctx.session.flush()
    ctx.session.add(DocumentDataRevision(request_id='history', payload_hash='hash', source_document_id=ctx.doc.id,
        data_kind='finance', record_id=item.id, actor_id=ctx.actor.id, before_data={'dedupe_key': source['dedupe_key']}, after_data={'dedupe_key': 'new-key'}))
    ctx.session.commit()
    result = report(ctx.session, ctx.snapshot, '2026-09', 'finance')
    assert len(result['rows']) == 4
    assert result['rows'][3][7] == 7000
    assert result['summary']['系統收入'] == 7000


def test_duplicate_source_transactions_are_not_matched_twice(ctx):
    duplicate = dict(ctx.data['finance'][3], source='duplicate')
    source_data = {**ctx.data, 'finance': [*ctx.data['finance'], duplicate]}
    ctx.snapshot.data = source_data
    source = ctx.data['finance'][3]
    ctx.session.add(FinanceEntry(entry_date=date(2026, 9, 4), entry_type='收入', category='工程款', amount=6000,
                                income_amount=6000, dedupe_key=source['dedupe_key']))
    ctx.session.commit()
    result = report(ctx.session, ctx.snapshot, '2026-09', 'finance')
    assert len(result['rows']) == 6
    assert result['rows'][3][7] is None and result['rows'][4][7] is None
    assert result['summary']['系統收入'] == 6000
    assert sum(row[7] or 0 for row in result['rows']) == 6000


def test_roster_taiwan_month_boundary_and_only_approved_leave(ctx):
    ctx.session.add_all([
        AttendanceEvent(employee_id=ctx.employee.id, event_type='上班打卡', happened_at=datetime(2026, 8, 31, 16, 30)),
        AttendanceEvent(employee_id=ctx.employee.id, event_type='下班打卡', happened_at=datetime(2026, 9, 30, 16, 0)),
        LeaveRequest(employee_id=ctx.employee.id, leave_type='病假', start_date=date(2026, 9, 2), end_date=date(2026, 9, 3), reason='病假', status=LeaveStatus.approved),
        LeaveRequest(employee_id=ctx.employee.id, leave_type='事假', start_date=date(2026, 9, 4), end_date=date(2026, 9, 4), reason='待審', status=LeaveStatus.pending),
    ])
    ctx.session.commit()
    data = report(ctx.session, ctx.snapshot, '2026-09', 'roster')
    assert data['summary']['系統打卡事件'] == 1
    assert data['rows'][1][3] == '上班打卡'
    assert data['rows'][1][4] == '核准病假'
    assert data['rows'][1][6] == ''


def test_contact_live_values_win_and_employees_are_separate(ctx):
    ctx.session.add(BusinessContact(normalized_key='tax:12345678', name='測試公司', tax_id='12345678', mobile='0999999999'))
    ctx.session.commit()
    data = report(ctx.session, ctx.snapshot, '2026-09', 'contacts')
    assert data['rows'][0][7] == '0999999999'
    assert len(data['rows']) == 3


def test_mapping_never_matches_small_excel_code_to_random_employee(ctx):
    mapping, _ = employee_matches(ctx.data, [ctx.actor])
    assert mapping['林育弘'] is None
    mapping, _ = employee_matches(ctx.data, [ctx.actor], {'林育弘': 'ADMIN001'})
    assert mapping['林育弘'] == 'ADMIN001'


def import_payload(ctx, **changes):
    return {'document_id': ctx.doc.id, 'expected_content_sha256': ctx.data['content_sha256'],
            'employee_mapping': {'林育弘': 'EMP001'}, **changes}


def test_import_idempotency_version_and_business_data_preserved(ctx):
    before = len(ctx.session.exec(select(ManagementWorkbookSnapshot)).all())
    response = ctx.client.post('/api/management-reports/import', json=import_payload(ctx))
    assert response.status_code == 200, response.text
    assert response.json()['backup']['status'] == 'saved'
    assert len(ctx.session.exec(select(ManagementWorkbookSnapshot)).all()) == before + 1
    replay = ctx.client.post('/api/management-reports/import', json=import_payload(ctx))
    assert replay.json()['replayed']
    assert len(ctx.session.exec(select(AdminAuditLog)).all()) == 1
    assert not ctx.session.exec(select(FinanceEntry)).all()
    assert not ctx.session.exec(select(LeaveRequest)).all()
    changed = ctx.client.post('/api/management-reports/import', json=import_payload(ctx, employee_mapping={'林育弘': None}))
    assert changed.status_code == 200
    snapshot = ctx.session.get(ManagementWorkbookSnapshot, changed.json()['snapshot_id'])
    assert snapshot.employee_mapping['林育弘'] is None


def test_changed_source_invalid_mapping_and_backup_failure(ctx):
    assert ctx.client.post('/api/management-reports/import', json=import_payload(ctx, expected_content_sha256='0' * 64)).status_code == 409
    assert ctx.client.post('/api/management-reports/import', json=import_payload(ctx, employee_mapping={'林育弘': 'fake'})).status_code == 422
    ctx.backup.side_effect = RuntimeError('do-not-disclose')
    result = ctx.client.post('/api/management-reports/import', json=import_payload(ctx))
    assert result.status_code == 200 and result.json()['saved']
    assert result.json()['backup']['status'] == 'failed'
    assert 'do-not-disclose' not in result.text


@pytest.mark.parametrize('month', ['2026-00', '2026-13', '2026-1', '2026-09-01'])
def test_month_validation(ctx, month):
    assert ctx.client.get('/api/management-reports/report', params={'month': month}).status_code == 422


def test_role_and_unauthenticated_protection(ctx):
    ctx.actor.role = Role.employee
    for path in ['/sources', '/preview?document_id=1', '/report?month=2026-09', '/export?month=2026-09']:
        assert ctx.client.get('/api/management-reports' + path).status_code == 403
    assert ctx.client.post('/api/management-reports/import', json=import_payload(ctx)).status_code == 403
    app.dependency_overrides.pop(get_current_actor)
    assert ctx.client.get('/api/management-reports/sources').status_code == 401


def test_paging_csv_full_rows_and_formula_injection(ctx):
    ctx.session.add(FinanceEntry(entry_date=date(2026, 9, 5), entry_type='收入', category='工程款', summary='=danger()', amount=1))
    ctx.session.commit()
    partial = ctx.client.get('/api/management-reports/report', params={'month': '2026-09', 'kind': 'finance', 'snapshot_id': ctx.snapshot.id, 'limit': 1})
    assert len(partial.json()['rows']) == 1 and partial.json()['total'] == 5
    response = ctx.client.get('/api/management-reports/export', params={'month': '2026-09', 'snapshot_id': ctx.snapshot.id})
    assert response.status_code == 200 and response.content.startswith(b'\xef\xbb\xbf')
    assert "'=danger()" in response.content.decode('utf-8-sig')
    assert response.headers['cache-control'] == 'no-store'
    assert len(response.content.decode('utf-8-sig').splitlines()) == 6
