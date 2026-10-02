from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from uuid import uuid4

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import (
    AdminAuditLog, BusinessContact, DocumentDataExport, DocumentDataRevision,
    Employee, FinanceEntry, ManagedDocument, Role,
)
from app.routes import document_data as routes
from app.routes.finance_imports import _import_contacts, _import_finance
from app.services.document_data import digest, snapshot
from app.services.finance_import import analyze_finance_workbook
from test_finance_imports import _annual_workbook_bytes


@pytest.fixture
def ctx():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    SQLModel.metadata.create_all(engine)
    session = Session(engine)
    actor = Employee(employee_code='ADMIN001', name='測試管理員', bind_token='test', role=Role.admin)
    document = ManagedDocument(category='通訊錄與年度管理', title='測試通訊錄', original_file_name='sample.xlsx',
                               stored_file_name='original.xlsx', drive_file_id='original', drive_folder_id='folder',
                               drive_url='https://drive.example/original')
    session.add_all([actor, document])
    session.commit()
    analysis = analyze_finance_workbook(_annual_workbook_bytes(), document.original_file_name)
    _import_finance(session, document=document, batch_id=None, records=analysis['finance_candidates'])
    _import_contacts(session, document=document, batch_id=None, contacts=analysis['contact_candidates'])
    session.commit()
    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[get_current_actor] = lambda: actor
    backup = AsyncMock(return_value={'status': 'saved'})
    upload = AsyncMock(return_value=SimpleNamespace(file_name='export.xlsx', file_id='new-version',
                                                    folder_id='folder', file_url='https://drive.example/new'))
    with patch.object(routes.drive, 'backup_database', backup), patch.object(routes.drive, 'upload_management_document', upload):
        yield SimpleNamespace(session=session, actor=actor, document=document, client=TestClient(app),
                              backup=backup, upload=upload, analysis=analysis)
    app.dependency_overrides.clear()
    session.close()
    engine.dispose()


def listed(ctx, kind='finance'):
    response = ctx.client.get(f'/api/document-data/{ctx.document.id}', params={'kind': kind})
    assert response.status_code == 200, response.text
    return response.json()


def payload(ctx, kind='finance', **changes):
    row = listed(ctx, kind)['rows'][0]
    values = {key: value if value is not None else '' for key, value in row['values'].items()}
    if kind == 'finance':
        values['invoice_date'] = row['values']['invoice_date']
    values.update(changes)
    return {'kind': kind, 'record_id': row['id'], 'expected_revision': row['revision'],
            'request_id': str(uuid4()), 'values': values}


def save(ctx, data):
    return ctx.client.post(f'/api/document-data/{ctx.document.id}/save', json=data)


def test_save_history_provenance_and_retry(ctx):
    data = payload(ctx, summary='人工修正', income_amount=12000)
    before = snapshot(ctx.session.get(FinanceEntry, data['record_id']))
    response = save(ctx, data)
    assert response.status_code == 200, response.text
    assert response.json()['backup']['status'] == 'saved'
    row = ctx.session.get(FinanceEntry, data['record_id'])
    assert row.amount == row.income_amount == 12000
    assert row.source_row == before['source_row']
    assert row.source_document_id == ctx.document.id
    revision = ctx.session.exec(select(DocumentDataRevision)).one()
    assert revision.before_data == before
    assert revision.after_data['summary'] == '人工修正'
    assert len(ctx.session.exec(select(AdminAuditLog)).all()) == 1
    replay = save(ctx, data)
    assert replay.status_code == 200 and replay.json()['replayed']
    assert len(ctx.session.exec(select(DocumentDataRevision)).all()) == 1
    data['values']['summary'] = '另一筆'
    assert save(ctx, data).status_code == 409


def test_stale_revision_never_overwrites(ctx):
    first = payload(ctx, summary='最新')
    stale = payload(ctx, summary='過期')
    assert save(ctx, first).status_code == 200
    assert save(ctx, stale).status_code == 409
    assert ctx.session.get(FinanceEntry, first['record_id']).summary == '最新'
    assert len(ctx.session.exec(select(DocumentDataRevision)).all()) == 1


@pytest.mark.parametrize('changes', [
    {'source_document_id': 5}, {'amount': 1}, {'entry_date': 'bad'},
    {'income_amount': -1}, {'income_amount': 0}, {'expense_amount': 200},
    {'income_amount': 1.234}, {'category': ' '}, {'income_amount': 'Infinity'},
])
def test_finance_invalid_fields_do_not_write(ctx, changes):
    data = payload(ctx, **changes)
    assert save(ctx, data).status_code == 422
    assert not ctx.session.exec(select(DocumentDataRevision)).all()
    assert ctx.session.get(FinanceEntry, data['record_id']).income_amount == 10000


def test_new_contact_unique_and_no_employee_changes(ctx):
    data = {'kind': 'contacts', 'request_id': str(uuid4()), 'values': {'name': '新供應商', 'mobile': '0912000000'}}
    result = save(ctx, data)
    assert result.status_code == 200, result.text
    assert save(ctx, data).json()['replayed']
    assert len(ctx.session.exec(select(BusinessContact)).all()) == 2
    data['request_id'] = str(uuid4())
    assert save(ctx, data).status_code == 409
    assert len(ctx.session.exec(select(Employee)).all()) == 1


def test_backup_failure_is_separate_from_committed_save(ctx):
    ctx.backup.side_effect = RuntimeError('secret-not-for-client')
    response = save(ctx, payload(ctx, summary='仍已保存'))
    assert response.status_code == 200
    assert response.json()['saved'] and response.json()['backup']['status'] == 'failed'
    assert 'secret-not-for-client' not in response.text
    assert listed(ctx)['rows'][0]['values']['summary'] == '仍已保存'


def test_other_document_cannot_edit_record(ctx):
    data = payload(ctx, summary='錯誤来源')
    other = ManagedDocument(category='收支明細', title='另一檔', original_file_name='other.xlsx',
                            stored_file_name='other.xlsx', drive_file_id='other', drive_folder_id='folder', drive_url='https://drive.example/other')
    ctx.session.add(other)
    ctx.session.commit()
    response = ctx.client.post(f'/api/document-data/{other.id}/save', json=data)
    assert response.status_code == 404
    assert not ctx.session.exec(select(DocumentDataRevision)).all()


def test_private_export_new_version_and_literal_strings(ctx):
    assert save(ctx, payload(ctx, summary='=HYPERLINK("unsafe")')).status_code == 200
    response = ctx.client.post(f'/api/document-data/{ctx.document.id}/export', json={
        'kind': 'finance', 'expected_dataset': listed(ctx)['dataset_hash'],
    })
    assert response.status_code == 200, response.text
    args = ctx.upload.call_args.kwargs
    assert args['public_share'] is False
    assert '線上資料匯出' in args['file_name'] and args['file_name'] != 'sample.xlsx'
    workbook = openpyxl.load_workbook(BytesIO(args['content']))
    assert workbook.active['D2'].value == '=HYPERLINK("unsafe")'
    assert workbook.active['D2'].data_type == 's'
    assert len(ctx.session.exec(select(DocumentDataExport)).all()) == 1
    assert ctx.session.get(ManagedDocument, ctx.document.id).drive_file_id == 'original'
    assert len(ctx.session.exec(select(ManagedDocument)).all()) == 2


def test_stale_export_and_upload_failure_leave_original_intact(ctx):
    old = listed(ctx)['dataset_hash']
    assert save(ctx, payload(ctx, summary='新的')).status_code == 200
    url = f'/api/document-data/{ctx.document.id}/export'
    assert ctx.client.post(url, json={'kind': 'finance', 'expected_dataset': old}).status_code == 409
    ctx.upload.assert_not_called()
    ctx.upload.side_effect = RuntimeError('private token')
    response = ctx.client.post(url, json={'kind': 'finance', 'expected_dataset': listed(ctx)['dataset_hash']})
    assert response.status_code == 502 and 'private token' not in response.text
    assert not ctx.session.exec(select(DocumentDataExport)).all()
    assert len(ctx.session.exec(select(ManagedDocument)).all()) == 1


def test_reimport_old_finance_and_contacts_respects_manual_changes(ctx):
    assert save(ctx, payload(ctx, summary='修正摘要', income_amount=12000)).status_code == 200
    assert save(ctx, payload(ctx, 'contacts', name='新公司名稱', tax_id='', phone='')).status_code == 200
    assert _import_finance(ctx.session, document=ctx.document, batch_id=None,
                          records=ctx.analysis['finance_candidates']) == (0, 2)
    assert _import_contacts(ctx.session, document=ctx.document, batch_id=None,
                           contacts=ctx.analysis['contact_candidates']) == (0, 0, 2)
    ctx.session.commit()
    assert len(ctx.session.exec(select(FinanceEntry)).all()) == 2
    contact = ctx.session.exec(select(BusinessContact)).one()
    assert contact.name == '新公司名稱' and not contact.phone and not contact.tax_id


def test_filters_pagination_and_worker_denied(ctx):
    result = ctx.client.get(f'/api/document-data/{ctx.document.id}', params={'keyword': '水泥', 'limit': 1}).json()
    assert result['total'] == 1 and result['dataset_total'] == 2
    assert result['rows'][0]['values']['summary'] == '水泥'
    ctx.actor.role = Role.employee
    assert ctx.client.get(f'/api/document-data/{ctx.document.id}').status_code == 403
    assert save(ctx, {'kind': 'contacts', 'request_id': str(uuid4()), 'values': {'name': '無權'}}).status_code == 403


def test_unauthenticated_and_wrong_category_denied(ctx):
    ctx.document.category = '推高機計價'
    ctx.session.commit()
    assert ctx.client.get(f'/api/document-data/{ctx.document.id}').status_code == 400
    app.dependency_overrides.pop(get_current_actor)
    assert ctx.client.get(f'/api/document-data/{ctx.document.id}').status_code == 401


def test_empty_export_rejected(ctx):
    contact = ctx.session.exec(select(BusinessContact)).one()
    ctx.session.delete(contact)
    ctx.session.commit()
    response = ctx.client.post(f'/api/document-data/{ctx.document.id}/export', json={
        'kind': 'contacts', 'expected_dataset': listed(ctx, 'contacts')['dataset_hash'],
    })
    assert response.status_code == 400
    ctx.upload.assert_not_called()


def test_atomic_write_race_rolls_back_without_history(ctx):
    data = payload(ctx, summary='模擬併發')
    real_execute = ctx.session.execute
    def lost_race(statement, *args, **kwargs):
        if getattr(statement, 'is_update', False):
            assert len(list(statement._where_criteria)) > 5
            return SimpleNamespace(rowcount=0)
        return real_execute(statement, *args, **kwargs)
    with patch.object(ctx.session, 'execute', side_effect=lost_race):
        assert save(ctx, data).status_code == 409
    assert not ctx.session.exec(select(DocumentDataRevision)).all()
    assert ctx.session.get(FinanceEntry, data['record_id']).summary == '6月工程'


def test_unchanged_save_and_backup_retry(ctx):
    data = payload(ctx)
    response = save(ctx, data)
    assert response.status_code == 200 and response.json()['unchanged']
    assert not ctx.session.exec(select(DocumentDataRevision)).all()
    assert ctx.client.post(f'/api/document-data/{ctx.document.id}/backup').json()['backup']['status'] == 'saved'


def test_old_finance_identity_cannot_be_inserted_as_new(ctx):
    original = payload(ctx)
    assert save(ctx, payload(ctx, summary='修改過')).status_code == 200
    original.pop('record_id')
    original.pop('expected_revision')
    assert save(ctx, original).status_code == 409
