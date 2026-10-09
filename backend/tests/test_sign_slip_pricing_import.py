"""簽單每台工時匯入：8+4 應為 9,200＋加班 1,000，不能把兩台都估成 8 小時。"""

from datetime import date
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine
from unittest.mock import AsyncMock, patch

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import Employee, ManagedDocument, Role, SignSlipRecord, WorkAssignment, Worksite
from app.routes import admin as admin_routes
from app.services.cost_workbook import (
    billable_hours_from_texts,
    money_differs,
    normal_day_amount,
    parse_vehicle_hour_segments,
    pooled_normal_amount,
    price_normal_hours,
    read_pricing_parameters,
)


RATES = (800, 6000, 1000, 1000, 800)


def _normal_formula(row: int) -> str:
    return (
        f'=IF(B{row}="","",IF(B{row}<8,B{row}*參數!$B$22,'
        f'INT(B{row}/8)*參數!$B$23+MOD(B{row},8)*參數!$B$22))'
    )


def _overtime_formula(row: int) -> str:
    return f'=IF(D{row}="","",D{row}*參數!$B$24)'


def _workbook_bytes(prefilled: dict[date, tuple[float, float]] | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    params = workbook.create_sheet("參數")
    params["B4"], params["B5"] = 0.05, 0.006
    params["B21"] = "53"
    params["B22"], params["B23"], params["B24"], params["B25"], params["B26"] = RATES
    sheet = workbook.create_sheet("11510-53")
    prefilled = prefilled or {}
    for offset, day_number in enumerate(range(1, 9)):
        row = 5 + offset
        day = date(2026, 10, day_number)
        sheet.cell(row=row, column=1, value=day)
        sheet.cell(row=row, column=3, value=_normal_formula(row))
        sheet.cell(row=row, column=5, value=_overtime_formula(row))
        sheet.cell(row=row, column=7, value=f'=IF(F{row}="","",F{row}*參數!$B$26)')
        if day in prefilled:
            normal, overtime = prefilled[day]
            sheet.cell(row=row, column=2, value=normal)
            sheet.cell(row=row, column=4, value=overtime)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


# 文字取自系統印出的四張簽單。10/7 工作內容先寫「加班3小時」「加班1小時」，
# 計價必須採用最後「工時計算」那一行的加班 4 小時，不能被敘述裡的數字帶走。
SLIPS = {
    date(2026, 10, 5): (
        "明旺石材 三台車下料\n工時計算：3 台正常作業 8+8+8  = 24 小時　加班 0  小時　共計 24 小時",
        24, 0, 3, 18000,
    ),
    date(2026, 10, 6): (
        "九州下料 兩台車\n建成晚上加班一小時\n工時計算：2 台正常作業 8+4  = 12 小時　加班 1  小時　共計 13 小時",
        12, 1, 2, 10200,
    ),
    date(2026, 10, 7): (
        "上午兩台堆高機加班3小時　建成 06：00〜08：00\n育弘晚上加班1小時 17：00〜18：00\n"
        "工時計算：3 台正常作業 8+8+8  = 24 小時　加班 4  小時　共計 28 小時",
        24, 4, 3, 22000,
    ),
    date(2026, 10, 8): (
        "紘維水電 一台大車和小車下料\n工時計算：3 台正常作業 8+8+8  = 24 小時　加班 1  小時　共計 25 小時",
        24, 1, 3, 19000,
    ),
}


def test_vehicle_expression_prices_each_machine_not_the_pooled_shortcut_when_they_differ():
    parameters = read_pricing_parameters(_workbook_bytes(), "計價.xlsx")
    rates = parameters.rates_for("53")
    segments = parse_vehicle_hour_segments("工時計算：2 台正常作業 7+7 = 14 小時")
    assert segments == (7, 7)
    assert price_normal_hours(rates, 14, segments) == 11200
    assert pooled_normal_amount(rates, 14, segments) == 10800
    assert money_differs(11200, 10800)
    # 8+4 與合計 12 小時在這個費率下相同：一台日薪 6,000、一台 4×800。
    split = parse_vehicle_hour_segments(SLIPS[date(2026, 10, 6)][0])
    assert split == (8, 4)
    assert price_normal_hours(rates, 12, split) == 9200
    assert normal_day_amount(rates, 12) == 9200
    assert normal_day_amount(rates, 16) == 12000


def test_four_october_slips_match_printed_amounts():
    parameters = read_pricing_parameters(_workbook_bytes(), "計價.xlsx")
    rates = parameters.rates_for("53")
    for day, (text, normal, overtime, _count, total) in SLIPS.items():
        parsed = billable_hours_from_texts([{
            "work_content": text, "notes": None, "normal_hours": normal,
            "overtime_hours": overtime, "is_active": True,
        }])
        assert parsed["segments"] is not None
        assert parsed["normal_hours"] == normal
        assert parsed["overtime_hours"] == overtime
        amount = price_normal_hours(rates, parsed["normal_hours"], parsed["segments"])
        amount += overtime * rates.ot_rate
        assert amount == total, day


class _PricingClient:
    def setup_method(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.admin = Employee(employee_code="ADMIN001", name="林金谷", bind_token="admin", role=Role.admin)
        self.site = Worksite(code="53", name="齊裕53")
        self.session.add_all([self.admin, self.site])
        self.session.commit()
        self.document = ManagedDocument(
            category="推高機計價", title="推高機計價",
            original_file_name="三通工程行115年推高機計價.xlsm",
            stored_file_name="cost.xlsm", drive_file_id="cost-file",
            drive_folder_id="folder", drive_url="https://drive.example/original",
            content_type="application/vnd.ms-excel.sheet.macroEnabled.12",
            uploaded_by_id=self.admin.id,
        )
        self.session.add(self.document)
        self.session.commit()
        self.download = AsyncMock(return_value=_workbook_bytes())
        self.upload = AsyncMock(return_value="https://drive.example/edited")
        self.backup = AsyncMock(return_value={"status": "saved"})
        self.patches = [
            patch.object(admin_routes.google_drive_worklog_service, "download_file_bytes", self.download),
            patch.object(admin_routes.google_drive_worklog_service, "update_file_bytes", self.upload),
            patch.object(admin_routes.google_drive_worklog_service, "backup_database", self.backup),
        ]
        for item in self.patches:
            item.start()
        app.dependency_overrides[get_session] = lambda: self.session
        app.dependency_overrides[get_current_actor] = lambda: self.admin
        self.client = TestClient(app)

    def teardown_method(self):
        app.dependency_overrides.clear()
        for item in self.patches:
            item.stop()
        self.session.close()
        self.engine.dispose()

    def _add_slip(self, day: date):
        text, normal, overtime, count, _total = SLIPS[day]
        self.session.add(SignSlipRecord(
            slip_no=f"ST-{day.strftime('%Y%m%d')}-53",
            slip_date=day,
            worksite_id=self.site.id,
            site_code="53",
            location="齊裕53",
            work_content=text,
            forklift_count=count,
            normal_hours=normal,
            overtime_hours=overtime,
            total_hours=normal + overtime,
            is_active=True,
        ))

    def _add_two_vehicles(self, day: date):
        self.session.add(WorkAssignment(
            work_date=day, site_id=self.site.id, work_item="下料",
            vehicle="2.5噸", equipment="3.0噸", supervisor_id=self.admin.id,
        ))


class TestOctoberSignSlipImport(_PricingClient):
    def test_month_preview_uses_slip_segments_instead_of_two_full_days(self):
        self.download.return_value = _workbook_bytes({date(2026, 10, 6): (12, 1)})
        for day in SLIPS:
            self._add_slip(day)
        self._add_two_vehicles(date(2026, 10, 6))
        self.session.commit()

        response = self.client.post("/api/cost-hour-imports/month-preview", json={
            "document_id": self.document.id, "year": 2026, "month": 10,
        })
        assert response.status_code == 200, response.text
        days = {day["date"]: day for day in response.json()["labels"][0]["days"] if day["date"] >= "2026-10-05"}
        october_6 = days["2026-10-06"]
        assert october_6["journal_hours_source"] == "sign_slip"
        assert october_6["vehicle_hours"] == [8, 4]
        assert october_6["proposed_hours"]["normal_hours"] == 12
        assert october_6["proposed_hours"]["overtime_hours"] == 1
        assert october_6["action"] == "skipped_exists"
        assert october_6["amount_changed"] is False
        assert days["2026-10-05"]["proposed_hours"] == {
            "normal_hours": 24, "overtime_hours": 0, "support_hours": 0,
        }
        assert days["2026-10-07"]["proposed_hours"]["overtime_hours"] == 4
        assert days["2026-10-08"]["proposed_hours"]["overtime_hours"] == 1

    def test_blank_days_import_slip_amounts_for_all_four_dates(self):
        for day in SLIPS:
            self._add_slip(day)
        self._add_two_vehicles(date(2026, 10, 6))
        self.session.commit()
        response = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 10,
        })
        assert response.status_code == 200, response.text
        written = self.upload.await_args.kwargs["content"]
        workbook = openpyxl.load_workbook(BytesIO(written), data_only=False)
        sheet = workbook["11510-53"]
        cached = openpyxl.load_workbook(BytesIO(written), data_only=True)["11510-53"]
        expected = {
            date(2026, 10, 5): (24, 0, 18000),
            date(2026, 10, 6): (12, 1, 10200),
            date(2026, 10, 7): (24, 4, 22000),
            date(2026, 10, 8): (24, 1, 19000),
        }
        for day, (normal, overtime, total) in expected.items():
            row = 5 + (day.day - 1)
            assert sheet[f"B{row}"].value == normal
            assert sheet[f"D{row}"].value == overtime
            assert cached[f"C{row}"].value + cached[f"E{row}"].value == total

    def test_two_full_days_do_not_replace_correct_amount_without_confirmation(self):
        self.download.return_value = _workbook_bytes({date(2026, 10, 6): (12, 1)})
        self._add_two_vehicles(date(2026, 10, 6))
        self.session.commit()
        preview = self.client.post("/api/cost-hour-imports/preview", json={
            "document_id": self.document.id,
            "worksite_id": self.site.id,
            "work_date": "2026-10-06",
            "target_label": "53",
            "normal_hours": 16,
            "overtime_hours": 0,
            "support_hours": 0,
        })
        assert preview.status_code == 200, preview.text
        body = preview.json()
        assert body["pricing"]["existing_amount"]["total"] == 10200
        assert body["pricing"]["incoming_amount"]["total"] == 12000
        assert body["requires_amount_confirmation"] is True

        blocked = self.client.post("/api/cost-hour-imports", json={
            "document_id": self.document.id,
            "worksite_id": self.site.id,
            "work_date": "2026-10-06",
            "target_label": "53",
            "normal_hours": 16,
            "overtime_hours": 0,
            "support_hours": 0,
        })
        assert blocked.status_code == 409
        self.upload.assert_not_awaited()

        month = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 10, "overwrite": True,
        })
        assert month.status_code == 409
        self.upload.assert_not_awaited()

    def test_confirmed_slip_repairs_sixteen_hour_cell_back_to_sign_slip(self):
        self.download.return_value = _workbook_bytes({date(2026, 10, 6): (16, 0)})
        self._add_slip(date(2026, 10, 6))
        self._add_two_vehicles(date(2026, 10, 6))
        self.session.commit()
        blocked = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 10, "overwrite": True,
        })
        assert blocked.status_code == 409
        self.upload.assert_not_awaited()

        repaired = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 10,
            "overwrite": True, "confirm_amount_changes": True,
        })
        assert repaired.status_code == 200, repaired.text
        written = self.upload.await_args.kwargs["content"]
        sheet = openpyxl.load_workbook(BytesIO(written), data_only=False)["11510-53"]
        cached = openpyxl.load_workbook(BytesIO(written), data_only=True)["11510-53"]
        assert sheet["B10"].value == 12
        assert sheet["D10"].value == 1
        assert cached["C10"].value == 9200
        assert cached["E10"].value == 1000

    def test_slip_check_lists_bad_import_without_writing(self):
        self.download.return_value = _workbook_bytes({date(2026, 10, 6): (16, 0)})
        self._add_slip(date(2026, 10, 6))
        self._add_slip(date(2026, 10, 5))
        self.session.commit()
        # 先用確認寫入一筆錯誤工時紀錄，模擬過去已匯入的 16 小時。
        created = self.client.post("/api/cost-hour-imports", json={
            "document_id": self.document.id,
            "worksite_id": self.site.id,
            "work_date": "2026-10-06",
            "target_label": "53",
            "normal_hours": 16,
            "overtime_hours": 0,
            "confirm_amount_change": True,
        })
        assert created.status_code == 200, created.text
        self.upload.reset_mock()

        report = self.client.get("/api/cost-hour-imports/slip-check", params={
            "document_id": self.document.id, "year": 2026, "month": 10,
        })
        assert report.status_code == 200, report.text
        body = report.json()
        assert body["read_only"] is True
        assert body["changes_amounts"] is False
        flagged = {day["date"]: day for day in body["days"]}
        assert "2026-10-06" in flagged
        assert flagged["2026-10-06"]["vehicle_hours"] == [8, 4]
        assert flagged["2026-10-06"]["workbook"]["current_total"] == 12000
        assert flagged["2026-10-06"]["workbook"]["sign_slip_total"] == 10200
        assert "2026-10-05" not in flagged
        self.upload.assert_not_awaited()
