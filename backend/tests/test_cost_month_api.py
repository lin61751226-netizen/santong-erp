from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from unittest.mock import AsyncMock, patch

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import (
    AdminAuditLog, Employee, Forklift, ForkliftInspection, ManagedDocument,
    Role, WorkAssignment, WorkHourImportLog, Worksite, WorksiteJournalHours,
)
from app.routes import admin as admin_routes


def _normal_formula(column: str, row: int) -> str:
    return (
        f'=IF(B{row}="","",IF(B{row}<8,B{row}*參數!${column}$22,'
        f'INT(B{row}/8)*參數!${column}$23+MOD(B{row},8)*參數!${column}$22))'
    )


def _overtime_formula(column: str, row: int, *, holiday: bool) -> str:
    rate_row = 25 if holiday else 24
    return f'=IF(D{row}="","",D{row}*參數!${column}${rate_row})'


def _support_formula(column: str, row: int) -> str:
    return f'=IF(F{row}="","",F{row}*參數!${column}$26)'


def _cost_workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    params = workbook.create_sheet("參數")
    params["B4"], params["B5"] = 0.05, 0.006
    params["B21"], params["I21"] = "45", "新竹寶山1"
    params["B22"], params["B23"], params["B24"], params["B25"], params["B26"] = 800, 6000, 1000, 1200, 800
    params["B27"], params["B28"] = 18900, 1
    params["I22"], params["I23"], params["I24"], params["I25"], params["I26"] = 1000, 8000, 1100, 1200, 800
    params["I27"], params["I28"] = 0, 9

    days = [date(2026, 9, d) for d in range(1, 4)]
    for label, column in (("45", "B"), ("新竹寶山1", "I")):
        sheet = workbook.create_sheet(f"11509-{label}")
        for index, day in enumerate(days, start=5):
            sheet.cell(row=index, column=1, value=day)
            sheet.cell(row=index, column=3, value=_normal_formula(column, index))
            sheet.cell(row=index, column=5, value=_overtime_formula(column, index, holiday=day == date(2026, 9, 2)))
            sheet.cell(row=index, column=7, value=_support_formula(column, index))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class TestCostMonthApi:
    def setup_method(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.admin = Employee(employee_code="ADMIN001", name="林金谷", bind_token="admin", role=Role.admin)
        self.driver = Employee(employee_code="D001", name="測試司機", bind_token="d1")
        self.site_45 = Worksite(code="45", name="示範工地")
        self.session.add_all([self.admin, self.driver, self.site_45])
        self.session.commit()

        # 9/1：派工 1 台 2.5T、點檢 1 台 3.0T → 共 2 台、正常 16 小時
        self.session.add(WorkAssignment(
            work_date=date(2026, 9, 1), site_id=self.site_45.id,
            work_item="堆高機卸料", vehicle="2.5噸", supervisor_id=self.admin.id,
        ))
        forklift_3t = Forklift(forklift_code="3T-1", model="自排 3.0噸柴油車", current_site_id=self.site_45.id)
        self.session.add(forklift_3t)
        self.session.commit()
        self.session.add(ForkliftInspection(
            forklift_id=forklift_3t.id, operator_id=self.driver.id, site_id=self.site_45.id,
            inspection_date=date(2026, 9, 1), inspection_items={}, all_passed=True,
            created_at=datetime(2026, 9, 1, 8, 0),
        ))
        # 9/2：派工 2 台 2.5T → 共 2 台、正常 16 小時
        self.session.add(WorkAssignment(
            work_date=date(2026, 9, 2), site_id=self.site_45.id,
            work_item="堆高機作業", vehicle="2.5噸 數量2", supervisor_id=self.admin.id,
        ))
        # 9/3：無派工、無點檢 → 不寫入
        self.document = ManagedDocument(
            category="推高機計價", title="推高機計價",
            original_file_name="三通工程行115年推高機計價.xlsm",
            stored_file_name="cost.xlsm", drive_file_id="cost-file-1",
            drive_folder_id="folder", drive_url="https://drive.example/original",
            content_type="application/vnd.ms-excel.sheet.macroEnabled.12",
            uploaded_by_id=self.admin.id,
        )
        self.session.add(self.document)
        self.session.commit()

        self.download = AsyncMock(return_value=_cost_workbook_bytes())
        self.upload = AsyncMock(return_value="https://drive.example/edited")
        self.backup = AsyncMock(return_value={"status": "saved"})
        self.patches = [
            patch.object(admin_routes.google_drive_worklog_service, "download_file_bytes", self.download),
            patch.object(admin_routes.google_drive_worklog_service, "update_file_bytes", self.upload),
            patch.object(admin_routes.google_drive_worklog_service, "backup_database", self.backup),
        ]
        for item in self.patches:
            item.start()

        app.dependency_overrides[get_session] = lambda: self.session
        app.dependency_overrides[get_current_actor] = lambda: self.admin
        self.client = TestClient(app)

    def teardown_method(self):
        app.dependency_overrides.clear()
        for item in self.patches:
            item.stop()
        self.session.close()
        self.engine.dispose()

    def test_month_preview_aggregates_journal_units_and_matches_label(self):
        response = self.client.post("/api/cost-hour-imports/month-preview", json={
            "document_id": self.document.id, "year": 2026, "month": 9, "overwrite": False,
        })
        assert response.status_code == 200, response.text
        plan = response.json()
        assert plan["period"] == "11509"
        assert plan["counts"]["write"] == 2
        assert plan["counts"]["empty"] == 1
        assert plan["counts"]["no_worksite"] == 3  # 新竹寶山1 三天皆無對應工地

        label_45 = next(item for item in plan["labels"] if item["label"] == "45")
        assert label_45["matched"] is True
        assert label_45["worksite_name"] == "示範工地"
        day1 = next(day for day in label_45["days"] if day["date"] == "2026-09-01")
        assert day1["units"] == {"twoPointFive": 1, "threePointZero": 1, "fourPointFive": 0}
        assert day1["forklift_count"] == 2
        assert day1["incoming_normal"] == 16
        assert day1["action"] == "write"
        assert day1["amount"]["normal_amount"] == 12000   # 16 小時：2 個日薪
        day2 = next(day for day in label_45["days"] if day["date"] == "2026-09-02")
        assert day2["units"]["twoPointFive"] == 2
        assert day2["incoming_normal"] == 16
        day3 = next(day for day in label_45["days"] if day["date"] == "2026-09-03")
        assert day3["action"] == "empty"

        baoshan = next(item for item in plan["labels"] if item["label"] == "新竹寶山1")
        assert baoshan["matched"] is False
        assert any("新竹寶山1" in warning for warning in plan["warnings"])

    def test_month_apply_writes_workbook_once_and_logs_each_day(self):
        response = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 9, "overwrite": False,
        })
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["written"] == 2
        assert result["empty"] == 1

        # 同一活頁簿整月只回寫一次
        self.upload.assert_awaited_once()
        written_bytes = self.upload.await_args.kwargs["content"]
        workbook = openpyxl.load_workbook(BytesIO(written_bytes), data_only=False)
        assert workbook["11509-45"]["B5"].value == 16
        assert workbook["11509-45"]["B6"].value == 16
        assert workbook["11509-45"]["B7"].value is None   # 無作業不寫入
        assert workbook["11509-45"]["C5"].value.startswith("=IF(")  # 公式保留

        logs = self.session.exec(select(WorkHourImportLog)).all()
        assert {(log.work_date, log.normal_hours, log.target_label) for log in logs} == {
            (date(2026, 9, 1), 16, "45"),
            (date(2026, 9, 2), 16, "45"),
        }
        audit = self.session.exec(
            select(AdminAuditLog).where(AdminAuditLog.action == "import_hours_month")
        ).all()
        assert len(audit) == 1
        self.backup.assert_awaited_once()

    def test_month_apply_without_journal_returns_400(self):
        response = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 8, "overwrite": False,
        })
        assert response.status_code == 400
        self.upload.assert_not_awaited()

    def test_month_preview_overrides_recalculate_amounts_and_invoice(self):
        response = self.client.post("/api/cost-hour-imports/month-preview", json={
            "document_id": self.document.id, "year": 2026, "month": 9, "overwrite": False,
            "overrides": [
                {"label": "45", "date": "2026-09-01", "normal_hours": 8, "overtime_hours": 2},
            ],
        })
        assert response.status_code == 200, response.text
        plan = response.json()
        label_45 = next(item for item in plan["labels"] if item["label"] == "45")
        day1 = next(day for day in label_45["days"] if day["date"] == "2026-09-01")
        assert day1["action"] == "write"
        assert day1["hours"]["normal_hours"] == 8       # 覆寫工作日誌推算的 16
        assert day1["hours"]["overtime_hours"] == 2
        assert day1["amount"]["normal_amount"] == 6000
        assert day1["amount"]["overtime_amount"] == 2000  # 9/1 週二平日費率
        # 請款：9/1 正常6000+加班2000，9/2 正常12000 → 未稅 20000、稅 1000、含稅 21000
        assert label_45["summary"]["invoice_untaxed"] == 20000
        assert label_45["summary"]["invoice_tax"] == 1000
        assert label_45["summary"]["invoice_taxed"] == 21000
        assert plan["invoice"]["taxed"] == 21000

    def test_saved_journal_hours_flow_into_month_pricing(self):
        self.session.add(WorksiteJournalHours(
            work_date=date(2026, 9, 1), worksite_id=self.site_45.id,
            normal_hours=10, overtime_hours=3, support_hours=2,
            updated_by_id=self.admin.id,
        ))
        self.session.commit()

        response = self.client.post("/api/cost-hour-imports/month-preview", json={
            "document_id": self.document.id, "year": 2026, "month": 9, "overwrite": False,
        })
        assert response.status_code == 200, response.text
        label_45 = next(item for item in response.json()["labels"] if item["label"] == "45")
        day1 = next(day for day in label_45["days"] if day["date"] == "2026-09-01")
        assert day1["journal_hours_source"] == "saved"
        assert day1["incoming_normal"] == 10
        assert day1["hours"] == {
            "normal_hours": 10,
            "overtime_hours": 3,
            "support_hours": 2,
        }

    def test_month_apply_with_overrides_writes_overtime(self):
        response = self.client.post("/api/cost-hour-imports/month", json={
            "document_id": self.document.id, "year": 2026, "month": 9, "overwrite": False,
            "overrides": [
                {"label": "45", "date": "2026-09-01", "overtime_hours": 2},
            ],
        })
        assert response.status_code == 200, response.text
        assert response.json()["written"] == 2
        written_bytes = self.upload.await_args.kwargs["content"]
        workbook = openpyxl.load_workbook(BytesIO(written_bytes), data_only=False)
        assert workbook["11509-45"]["B5"].value == 16   # 工作日誌正常工時
        assert workbook["11509-45"]["D5"].value == 2    # 人工調整的加班工時
        cached = openpyxl.load_workbook(BytesIO(written_bytes), data_only=True)
        assert cached["11509-45"]["E5"].value == 2000   # 快取一併回填
