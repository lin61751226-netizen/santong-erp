from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock, patch

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.core.db import get_session
from app.deps import get_current_actor
from app.main import app
from app.models import BusinessContact, Employee, FinanceEntry, ManagedDocument, Role
from app.routes import finance_imports as finance_routes
from app.services.finance_import import analyze_finance_workbook


def _annual_workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    contacts = workbook.active
    contacts.title = "公司通訊錄"
    contacts.append([])
    contacts.append([])
    contacts.append(["序號", "姓名/公司", "統一編號", "部門", "職稱", "電話", "手機", "Email", "地址", "類別", "負責人", "備註"])
    contacts.append([1, "齊裕營造股份有限公司", "23344962", "", "財務", "", "", "", "高雄市", "業主", "吳小姐"])
    contacts.append([2, "齊裕營造股份有限公司", "23344962", "", "", "07-5522755", "", "", "", "合作夥伴", ""])
    contacts.append([3, "黃富俊(怪手)", "", "006合庫", "南勢角分行", "0720-765873340"])

    finance = workbook.create_sheet("115年06月_收支明細表")
    finance.append([])
    finance.append([])
    finance.append(["日期", "類型", "類別", "項目/摘要", "對象/公司", "付款方式", "憑證", "收入", "支出", "淨額", "累計", "經手人", "備註", "月份", "資金帳戶", "交易狀況", "憑證日期", "手續費", "營業稅", "憑證類型", "標籤"])
    finance.append([date(2026, 6, 1), "收款", "工程款收入", "6月工程", "齊裕營造", "銀行轉帳", "TX-001", 10000, 0, None, None, "秀蓉"])
    finance.append([date(2026, 6, 2), "收款", "工程款收入", "含手續費", "齊裕營造", "銀行轉帳", "TX-002", 5000, 20])
    finance.append([date(2026, 6, 3), "付款", "", "材料", "供應商", "現金", "", 0, 200])
    finance.append([date(2026, 6, 4), "付款", "材料費", "水泥", "武雄", "轉帳", "P-001", 0, 3000])
    finance.append([date(2026, 6, 5), "付款", "材料費", "空白金額", "武雄"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _liuhe_workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    contacts = workbook.active
    contacts.title = "公司通訊錄"
    contacts.append([])
    contacts.append([])
    contacts.append(["編號", "姓名/公司", "統一編號", "部門", "職稱", "公司電話", "分機", "手機", "Email", "地址", "類別", "負責人", "備註"])
    # Legacy data values remain in the older I:L positions.
    contacts.append([1, "文和採購", "12345678", "", "", "03-1234567", "", "0912345678", "桃園市", "供應商", "王小姐", "水泥"])
    finance = workbook.create_sheet("115年06月_收支明細表")
    finance.append([])
    finance.append([])
    finance.append(["日期", "收入/支出", "類別", "摘要", "對象/公司", "付款方式", "憑證號碼", "收入", "支出"])
    finance.append([date(2026, 6, 23), "收入", "工程款", "基座", "六和-中壢廠", "匯款", "L-001", 36640, 0])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class TestFinanceImport:
    def setup_method(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.admin = Employee(employee_code="ADMIN001", name="林金谷", bind_token="admin", role=Role.admin)
        self.document = ManagedDocument(
            category="通訊錄與年度管理", title="全年管理",
            original_file_name="三通公司通訊錄_收支明細表_堆高機拉_全年管理版.xlsx",
            stored_file_name="annual.xlsx", drive_file_id="annual-file", drive_folder_id="folder",
            drive_url="https://drive.example/annual", uploaded_by_id=1,
        )
        self.session.add_all([self.admin, self.document])
        self.session.commit()
        self.session.refresh(self.document)
        app.dependency_overrides[get_session] = lambda: self.session
        app.dependency_overrides[get_current_actor] = lambda: self.admin
        self.download = AsyncMock(return_value=_annual_workbook_bytes())
        self.backup = AsyncMock(return_value={"status": "saved"})
        self.patches = [
            patch.object(finance_routes.google_drive_worklog_service, "download_file_bytes", self.download),
            patch.object(finance_routes.google_drive_worklog_service, "backup_database", self.backup),
        ]
        for item in self.patches:
            item.start()
        self.client = TestClient(app)

    def teardown_method(self):
        app.dependency_overrides.clear()
        for item in self.patches:
            item.stop()
        self.session.close()
        self.engine.dispose()

    def test_preview_then_commit_keeps_warnings_out_of_finance_ledger(self):
        preview = self.client.post("/api/finance-imports/preview", json={"document_id": self.document.id})
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["summary"] == {"contact_candidates": 2, "finance_candidates": 2, "warning_rows": 4}

        commit = self.client.post("/api/finance-imports/commit", json={
            "document_id": self.document.id,
            "expected_content_sha256": payload["content_sha256"],
            "import_contacts": True,
            "import_finance": True,
        })
        assert commit.status_code == 201, commit.text
        result = commit.json()
        assert result["batch"]["finance_imported"] == 2
        assert result["batch"]["finance_pending_review"] == 4
        assert result["batch"]["contact_created"] == 1
        assert result["backup"]["status"] == "saved"
        self.backup.assert_awaited_once()

        entries = self.session.exec(select(FinanceEntry).order_by(FinanceEntry.entry_date)).all()
        assert [(entry.entry_type, entry.amount, entry.summary) for entry in entries] == [
            ("收入", 10000, "6月工程"),
            ("支出", -3000, "水泥"),
        ]
        contact = self.session.exec(select(BusinessContact)).one()
        assert contact.name == "齊裕營造股份有限公司"
        assert contact.phone == "07-5522755"

        listed = self.client.get("/api/finance-imports/entries")
        assert listed.status_code == 200
        assert [row["summary"] for row in listed.json()] == ["水泥", "6月工程"]

        duplicate = self.client.post("/api/finance-imports/commit", json={
            "document_id": self.document.id,
            "expected_content_sha256": payload["content_sha256"],
        })
        assert duplicate.status_code == 409
        assert self.session.exec(select(FinanceEntry)).all().__len__() == 2

    def test_liuhe_legacy_contacts_and_project_name_are_recognized(self):
        analysis = analyze_finance_workbook(_liuhe_workbook_bytes(), "三通公司收支明細表_六和.xlsx")
        contact = analysis["contact_candidates"][0]
        record = analysis["finance_candidates"][0]
        assert contact["address"] == "桃園市"
        assert contact["category"] == "供應商"
        assert contact["contact_person"] == "王小姐"
        assert record["project_name"] == "六和"
        assert record["entry_type"] == "收入"
