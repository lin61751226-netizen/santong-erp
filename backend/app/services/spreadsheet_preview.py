from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from math import isfinite
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def preview_spreadsheet(content: bytes, file_name: str, sheet: str | None,
                        row_start: int, column_start: int) -> dict:
    """Read cached values only; never save, recalculate, or run workbook macros."""
    if Path(file_name).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("此格式請使用「開啟原始 Excel」閱讀；大字閱讀支援 .xlsx、.xlsm")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("Excel 解壓縮內容過大，請開啟原始 Excel 閱讀")
    except BadZipFile as exc:
        raise ValueError("無法讀取 Excel 格式，請確認原檔") from exc
    values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
        names = values.sheetnames
        if not names:
            raise ValueError("Excel 沒有可讀取的工作表")
        selected = sheet or names[0]
        if selected not in names:
            raise ValueError("找不到指定工作表，請重新選擇")
        ws, source = values[selected], formulas[selected]
        total_rows, total_columns = max(ws.max_row or 1, 1), max(ws.max_column or 1, 1)
        if row_start > total_rows or column_start > total_columns:
            raise ValueError("指定範圍超過工作表，請重新選擇工作表")
        row_end, column_end = min(row_start + 39, total_rows), min(column_start + 15, total_columns)
        bounds = dict(min_row=row_start, max_row=row_end, min_col=column_start, max_col=column_end)
        rows, has_uncached = [], False
        for number, (cached_row, formula_row) in enumerate(zip(ws.iter_rows(**bounds), source.iter_rows(**bounds)), row_start):
            cells = []
            for cell, original in zip(cached_row, formula_row):
                value = cell.value
                uncached = original.data_type == "f" and value is None
                has_uncached |= uncached
                numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                if uncached:
                    text = "（公式未快取）"
                elif value is None:
                    text = ""
                elif isinstance(value, (date, datetime, time)):
                    text = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
                elif isinstance(value, bool):
                    text = "TRUE" if value else "FALSE"
                elif numeric and not isfinite(value):
                    text = "（無效數值）"
                    numeric = False
                elif numeric:
                    # Percent formats scale the display, not the stored value.
                    text = f"{value * 100:,.2f}%" if "%" in cell.number_format else f"{value:,}"
                else:
                    text = str(value)
                cells.append({"text": text, "numeric": numeric, "uncached_formula": uncached})
            rows.append({"number": number, "cells": cells})
        return {"sheets": names, "sheet": selected, "total_rows": total_rows,
                "total_columns": total_columns, "row_start": row_start, "column_start": column_start,
                "columns": [get_column_letter(i) for i in range(column_start, column_end + 1)],
                "rows": rows, "has_uncached_formulas": has_uncached}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 內容無法預覽，請開啟原始 Excel 核對") from exc
    finally:
        if values is not None:
            values.close()
        if formulas is not None:
            formulas.close()
