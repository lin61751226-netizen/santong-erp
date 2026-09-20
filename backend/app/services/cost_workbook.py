"""推高機計價活頁簿的讀寫規則。

計價表的日期在 A 欄，正常／加班／支援工時分別在 B、D、F 欄。這個模組
只處理使用者明確確認的單一日期與單一標別，避免批次覆蓋既有計價紀錄。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl


SHEET_PATTERN = re.compile(r"^(?P<period>\d{5})-(?P<label>.+)$")
HOUR_COLUMNS = {"normal_hours": "B", "overtime_hours": "D", "support_hours": "F"}


class CostWorkbookError(ValueError):
    pass


@dataclass(frozen=True)
class CostSheetTarget:
    worksheet_name: str
    label: str
    row_number: int
    normal_hours: float | int | None
    overtime_hours: float | int | None
    support_hours: float | int | None


def roc_period(target_date: date) -> str:
    """將西元日期轉成計價表使用的民國年月，例如 2026-09 -> 11509。"""
    return f"{target_date.year - 1911:03d}{target_date.month:02d}"


def _load_workbook(content: bytes, file_name: str, *, read_only: bool):
    if not content:
        raise CostWorkbookError("計價檔內容為空白")
    return openpyxl.load_workbook(
        io.BytesIO(content),
        read_only=read_only,
        data_only=False,
        keep_vba=Path(file_name).suffix.lower() == ".xlsm",
    )


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _number(value: object) -> float | int | None:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _sheet_targets(workbook, target_date: date) -> list[CostSheetTarget]:
    period = roc_period(target_date)
    targets: list[CostSheetTarget] = []
    for worksheet in workbook.worksheets:
        match = SHEET_PATTERN.match(worksheet.title)
        if not match or match.group("period") != period:
            continue
        label = match.group("label")
        if label.startswith("小計") or label.startswith("總表"):
            continue
        for row_number in range(5, worksheet.max_row + 1):
            if _as_date(worksheet.cell(row=row_number, column=1).value) != target_date:
                continue
            targets.append(CostSheetTarget(
                worksheet_name=worksheet.title,
                label=label,
                row_number=row_number,
                normal_hours=_number(worksheet.cell(row=row_number, column=2).value),
                overtime_hours=_number(worksheet.cell(row=row_number, column=4).value),
                support_hours=_number(worksheet.cell(row=row_number, column=6).value),
            ))
            break
    return sorted(targets, key=lambda item: item.label)


def list_cost_sheet_targets(content: bytes, file_name: str, target_date: date) -> list[CostSheetTarget]:
    workbook = _load_workbook(content, file_name, read_only=True)
    try:
        targets = _sheet_targets(workbook, target_date)
    finally:
        workbook.close()
    if not targets:
        raise CostWorkbookError(f"計價表內找不到 {roc_period(target_date)} 月份或 {target_date.isoformat()} 的日期列")
    return targets


def find_cost_sheet_target(content: bytes, file_name: str, target_date: date, label: str) -> CostSheetTarget:
    normalized_label = label.strip()
    targets = list_cost_sheet_targets(content, file_name, target_date)
    for target in targets:
        if target.label == normalized_label:
            return target
    raise CostWorkbookError(f"找不到標別「{normalized_label}」的 {target_date.isoformat()} 計價列")


def import_cost_hours(
    content: bytes,
    file_name: str,
    target_date: date,
    label: str,
    *,
    normal_hours: float,
    overtime_hours: float,
    support_hours: float,
) -> tuple[bytes, CostSheetTarget]:
    """將單筆工時寫入指定日期／標別，保留活頁簿公式與 VBA。"""
    values = {
        "normal_hours": normal_hours,
        "overtime_hours": overtime_hours,
        "support_hours": support_hours,
    }
    if any(value < 0 for value in values.values()):
        raise CostWorkbookError("工時不可為負數")

    workbook = _load_workbook(content, file_name, read_only=False)
    try:
        target = find_cost_sheet_target(content, file_name, target_date, label)
        worksheet = workbook[target.worksheet_name]
        for field, column in HOUR_COLUMNS.items():
            value = values[field]
            worksheet[f"{column}{target.row_number}"] = int(value) if float(value).is_integer() else value
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), target
    finally:
        workbook.close()
