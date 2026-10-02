"""Business-field editing; never round-trip the original workbook or macros."""

import hashlib
import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models import BusinessContact, FinanceEntry


FIELDS = {
    "finance": [
        ("entry_date", "日期", "date"), ("entry_type", "收支類型", "select"),
        ("category", "類別", "text"), ("summary", "工作／交易摘要", "textarea"),
        ("vendor_name", "對象／公司", "text"), ("income_amount", "收入金額", "number"),
        ("expense_amount", "支出金額", "number"), ("project_name", "工地／專案", "text"),
        ("payment_method", "付款方式", "text"), ("payment_status", "收付款狀態", "text"),
        ("invoice_number", "憑證／發票號碼", "text"), ("invoice_date", "憑證／發票日期", "date"),
        ("handled_by", "經手人", "text"), ("account_name", "資金帳戶", "text"),
        ("voucher_type", "憑證類型", "text"), ("tag", "標籤", "text"), ("note", "備註", "textarea"),
    ],
    "contacts": [
        ("name", "姓名／公司", "text"), ("tax_id", "統一編號", "text"),
        ("category", "類別", "text"), ("department", "部門", "text"), ("title", "職稱", "text"),
        ("contact_person", "聯絡人", "text"), ("phone", "電話", "text"),
        ("mobile", "手機", "text"), ("email", "Email", "text"),
        ("address", "地址", "textarea"), ("note", "備註", "textarea"),
    ],
}
MODELS = {"finance": FinanceEntry, "contacts": BusinessContact}


def digest(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def snapshot(row) -> dict:
    return row.model_dump(mode="json")


def serialize(row, kind: str) -> dict:
    data = snapshot(row)
    return {
        "id": row.id, "revision": digest(data),
        "values": {key: data.get(key) for key, _, _ in FIELDS[kind]},
        "source_sheet": row.source_sheet, "source_row": row.source_row,
    }


def export_xlsx(rows: list, kind: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收支明細" if kind == "finance" else "公司通訊錄"
    fields = FIELDS[kind] + [("source_sheet", "來源工作表", "text"), ("source_row", "來源列", "number")]
    sheet.append([label for _, label, _ in fields])
    for row in rows:
        sheet.append([getattr(row, key) for key, _, _ in fields])
        # Treat user strings literally, including leading '='; no executable formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="285C4B")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
