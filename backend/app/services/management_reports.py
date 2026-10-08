"""Read source evidence and reconcile it with live data without changing operations."""

from __future__ import annotations

import calendar
import hashlib
import math
import re
from zipfile import ZipFile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import openpyxl
from sqlmodel import Session, select

from app.models import (
    AssignmentMember, AttendanceEvent, BusinessContact, DocumentDataRevision,
    Employee, FinanceEntry, LeaveRequest, LeaveStatus, MeetingRecord,
    WorkAssignment, Worksite,
)
from app.services.finance_import import _parse_contacts, finance_dedupe_key


SHEET_PATTERN = re.compile(r"^(\d{3})年(\d{2})月_(收支明細表|薪資表|排休表|行事曆記錄)$")
PAY_FIELDS = ["底薪", "職務津貼", "交通/伙食", "加班費", "獎金", "其他加給",
              "勞保", "健保", "勞退自提", "請假扣款", "其他扣款"]
KINDS = {"finance": "每月收支明細表", "payroll": "每月薪資表", "roster": "每月排休表",
         "calendar": "每月行事曆記錄", "annual": "年度總表", "categories": "年度類別統計",
         "contacts": "公司通訊錄"}
TZ = ZoneInfo("Asia/Taipei")


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def money_sum(values):
    values = list(values)
    if any(value is None for value in values):
        return None
    return float(sum((Decimal(str(value)) for value in values), Decimal(0)))


def date_text(value):
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else None


def local_time(value):
    return (value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value).astimezone(TZ)


def analyze_workbook(content: bytes):
    with ZipFile(BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 150 * 1024 * 1024:
            raise ValueError("Excel解壓後超過150MB，請拆分資料後再匯入")
    # Use normal mode: these supplied files omit worksheet dimensions in XML.
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    result = {kind: [] for kind in ("finance", "payroll", "roster", "calendar", "contacts")}
    result.update(warnings=[], source_notes=[], sheets=workbook.sheetnames, content_sha256=hashlib.sha256(content).hexdigest())
    try:
        contacts, warnings = _parse_contacts(workbook)
        for contact in contacts:
            result["contacts"].append({**contact, "source": f"{contact['source_sheet']}!{contact['source_row']}"})
        result["warnings"].extend(warnings)
        matched_sheets = 0
        for sheet in workbook:
            match = SHEET_PATTERN.fullmatch(sheet.title)
            if not match:
                continue
            matched_sheets += 1
            year, month = int(match[1]) + 1911, int(match[2])
            if not 1 <= month <= 12:
                raise ValueError("來源工作表月份不正確")
            period = f"{year:04d}-{month:02d}"
            def value(row, col):
                return sheet.cell(row, col).value
            def origin(row):
                return {"month": period, "source": f"{sheet.title}!{row}"}
            if match[3] == "收支明細表":
                for row in range(4, sheet.max_row + 1):
                    day = date_text(value(row, 1))
                    if not day:
                        continue
                    if day[:7] != period:
                        result["warnings"].append({"source": origin(row)["source"], "reason": "日期與工作表月份不同，依實際日期列報"})
                    income, expense = number(value(row, 8)), number(value(row, 9))
                    if income == expense == 0:
                        continue
                    typ, category = text(value(row, 2)), text(value(row, 3))
                    party = text(value(row, 5))
                    status = "待確認"
                    if "結餘" in party or "結餘" in text(value(row, 4)):
                        status = "結餘（非營收）"
                    elif typ == "轉帳":
                        status = "轉帳（非營收）"
                    elif category and income is not None and expense is not None and not (income and expense):
                        if (typ in {"收入", "收款"} and income > 0 and expense == 0) or (typ in {"支出", "付款"} and expense > 0 and income == 0):
                            status = "可對照"
                    item = {**origin(row), "month": day[:7], "date": day, "type": typ,
                            "category": category, "summary": text(value(row, 4)), "counterparty": party,
                            "payment_method": text(value(row, 6)), "voucher": text(value(row, 7)),
                            "income": income, "expense": expense, "handled_by": text(value(row, 12)),
                            "note": text(value(row, 13)), "account": text(value(row, 15)),
                            "status": status, "dedupe_key": None}
                    if status == "可對照":
                        item["dedupe_key"] = finance_dedupe_key({"entry_date": date.fromisoformat(day),
                            "entry_type": "收入" if typ in {"收入", "收款"} else "支出", "category": category,
                            "summary": item["summary"], "counterparty": party, "income_amount": income,
                            "expense_amount": expense, "voucher_number": item["voucher"], "payment_method": item["payment_method"]})
                    else:
                        result["warnings"].append({"source": item["source"], "reason": status})
                    result["finance"].append(item)
            elif match[3] == "薪資表":
                for row in range(6, sheet.max_row + 1):
                    name = text(value(row, 3))
                    if not name:
                        continue
                    inputs = [value(row, col) for col in range(6, 17)]
                    if not any(isinstance(item, (int, float)) and not isinstance(item, bool) for item in inputs) and not any(sheet.cell(row, col).data_type == "f" for col in range(6, 17)):
                        result["source_notes"].append({**origin(row), "name": name, "values": [text(value(row, col)) for col in range(1, 22)]})
                        result["warnings"].append({"source": origin(row)["source"], "reason": "薪資區文字註記（無金額），另存原註記，不列薪資人數/合計"})
                        continue
                    amounts = {key: number(value(row, col)) for col, key in enumerate(PAY_FIELDS, 6)}
                    gross = money_sum(amounts[key] for key in PAY_FIELDS[:6])
                    deductions = money_sum(amounts[key] for key in PAY_FIELDS[6:])
                    source_month = date_text(value(row, 1))
                    if source_month and source_month[:7] != period:
                        result["warnings"].append({"source": origin(row)["source"], "reason": "薪資列月份與工作表不同，未匯入"})
                        continue
                    item = {**origin(row), "name": name, "employee_code": text(value(row, 2)),
                            "department": text(value(row, 4)), "title": text(value(row, 5)),
                            **amounts, "gross": gross, "deductions": deductions,
                            "net": round(gross - deductions, 2) if gross is not None and deductions is not None else None,
                            "status": text(value(row, 20)), "note": text(value(row, 21))}
                    if item["net"] is None:
                        result["warnings"].append({"source": item["source"], "reason": "薪資輸入含文字／公式，金額待確認，不當成零"})
                    result["payroll"].append(item)
            elif match[3] == "排休表":
                days = calendar.monthrange(year, month)[1]
                for row in range(6, sheet.max_row + 1):
                    name = text(value(row, 2))
                    if not name:
                        continue
                    marks = {str(day): text(value(row, day + 4)) for day in range(1, days + 1) if text(value(row, day + 4))}
                    result["roster"].append({**origin(row), "name": name, "employee_code": text(value(row, 1)),
                                             "department": text(value(row, 3)), "title": text(value(row, 4)),
                                             "marks": marks, "note": text(value(row, 42))})
            else:
                for row in range(6, sheet.max_row + 1):
                    title = text(value(row, 12))
                    day = date_text(value(row, 8))
                    if title and day:
                        if day[:7] != period:
                            result["warnings"].append({"source": origin(row)["source"], "reason": "事項日期與工作表月份不同，依實際日期列報"})
                        result["calendar"].append({**origin(row), "month": day[:7], "date": day,
                            "title": title, "time": text(value(row, 10)), "type": text(value(row, 11)),
                            "location": text(value(row, 13)), "people": text(value(row, 14)),
                            "reminder": text(value(row, 15)), "status": text(value(row, 16)),
                            "completed_date": date_text(value(row, 17)), "note": text(value(row, 18))})
                    elif title:
                        result["warnings"].append({"source": origin(row)["source"], "reason": "行事曆事項缺少日期，未匯入"})
                start = date(year, month, 1)
                start -= timedelta(days=(start.weekday() + 1) % 7)
                for week in range(6):
                    for col in range(1, 8):
                        cell = sheet.cell(6 + week * 2, col)
                        title = text(cell.value)
                        day = start + timedelta(days=week * 7 + col - 1)
                        if title and cell.data_type != "f":
                            if day.month != month:
                                result["warnings"].append({"source": f"{sheet.title}!{cell.coordinate}", "reason": "月曆註記在本月以外，未匯入"})
                                continue
                            result["calendar"].append({"month": period, "date": day.isoformat(), "title": title,
                                "time": "", "type": "月曆註記", "location": "", "people": "", "status": "",
                                "note": "", "source": f"{sheet.title}!{cell.coordinate}"})
        if not matched_sheets:
            raise ValueError("找不到民國年月命名的管理工作表，請選擇全年管理版")
        key_counts = Counter(item["dedupe_key"] for item in result["finance"] if item["dedupe_key"])
        for item in result["finance"]:
            if item["dedupe_key"] and key_counts[item["dedupe_key"]] > 1:
                result["warnings"].append({"source": item["source"], "reason": "來源有相同交易多列，需人工對照，不重複配對正式帳"})
        return result
    finally:
        workbook.close()


def employee_matches(data, employees, overrides=None):
    overrides = overrides or {}
    names = sorted({item["name"] for kind in ("payroll", "roster") for item in data[kind]})
    by_code = {item.employee_code: item for item in employees}
    mapping, candidates = {}, []
    for name in names:
        codes = {item.get("employee_code") for kind in ("payroll", "roster") for item in data[kind] if item["name"] == name}
        exact = [item for item in employees if item.name.strip() == name]
        code_matches = [item for item in employees if item.employee_code in codes]
        hits = {item.id: item for item in [*exact, *code_matches]}
        selected = by_code.get(overrides[name]) if overrides.get(name) else next(iter(hits.values())) if len(hits) == 1 else None
        mapping[name] = selected.employee_code if selected else None
        candidates.append({"name": name, "source_codes": sorted(filter(None, codes)),
                           "employee_code": mapping[name], "status": "已對應" if selected else "待對應"})
    return mapping, candidates


def finance_values(row):
    income, expense = row.income_amount, row.expense_amount
    # Older manually-created entries only populated `amount`.
    if not income and not expense:
        income = abs(row.amount) if row.entry_type in {"收入", "收款"} else 0
        expense = abs(row.amount) if row.entry_type in {"支出", "付款"} else 0
    return income, expense


def report(session: Session, snapshot, month: str, kind: str):
    year, number_month = map(int, month.split("-"))
    first = date(year, number_month, 1)
    last = date(year, number_month, calendar.monthrange(year, number_month)[1])
    data = snapshot.data if snapshot else {key: [] for key in ("finance", "payroll", "roster", "calendar", "contacts")}
    mapping = snapshot.employee_mapping if snapshot else {}
    employees = session.exec(select(Employee).order_by(Employee.employee_code)).all()
    by_code = {row.employee_code: row for row in employees}
    finance = session.exec(select(FinanceEntry).where(FinanceEntry.entry_date >= date(year, 1, 1), FinanceEntry.entry_date < date(year + 1, 1, 1)).order_by(FinanceEntry.entry_date, FinanceEntry.id)).all()
    aliases = {}
    for item in finance:
        if item.dedupe_key:
            aliases[item.dedupe_key] = item
    by_id = {item.id: item for item in finance}
    for revision in session.exec(select(DocumentDataRevision).where(DocumentDataRevision.data_kind == "finance")).all():
        if revision.record_id in by_id:
            for values in (revision.before_data, revision.after_data):
                if values.get("dedupe_key"):
                    aliases[values["dedupe_key"]] = by_id[revision.record_id]
    source_rows = lambda key: [item for item in data[key] if item.get("month") == month]
    year_source_finance = [item for item in data["finance"] if item["month"].startswith(f"{year}-")]
    source_key_counts = Counter(item["dedupe_key"] for item in year_source_finance if item["dedupe_key"])
    summary = {}
    rows = []
    if kind == "finance":
        fields = ["日期", "類型", "類別", "摘要", "對象", "收入", "支出", "系統收入", "系統支出", "對照狀態", "來源"]
        matched = set()
        for item in source_rows("finance"):
            ambiguous = item["dedupe_key"] and source_key_counts[item["dedupe_key"]] > 1
            target = aliases.get(item["dedupe_key"]) if not ambiguous else None
            target_values = finance_values(target) if target else (None, None)
            if target:
                matched.add(target.id)
            rows.append([item["date"], item["type"], item["category"], item["summary"], item["counterparty"],
                         item["income"], item["expense"], *target_values,
                         "已對應系統（以系統值為準）" if target else "來源同交易多列，待人工對照" if ambiguous else f"Excel：{item['status']}，未入正式帳", item["source"]])
        for item in finance:
            if first <= item.entry_date <= last and item.id not in matched:
                rows.append([item.entry_date.isoformat(), item.entry_type, item.category, item.summary,
                             item.vendor_name, None, None, *finance_values(item), "系統已存", f"FinanceEntry #{item.id}"])
        summary = {"系統收入": money_sum(finance_values(item)[0] for item in finance if first <= item.entry_date <= last),
                   "系統支出": money_sum(finance_values(item)[1] for item in finance if first <= item.entry_date <= last),
                   "Excel原始收入（含結餘）": money_sum(item["income"] for item in source_rows("finance")),
                   "Excel原始支出（含轉帳）": money_sum(item["expense"] for item in source_rows("finance"))}
    elif kind == "payroll":
        fields = ["月份", "Excel姓名", "系統員工", "職稱", *PAY_FIELDS, "應領合計", "扣款合計", "實領薪資", "發薪狀態", "對應狀態", "備註", "來源"]
        for item in source_rows("payroll"):
            employee = by_code.get(mapping.get(item["name"]))
            rows.append([month, item["name"], f"{employee.employee_code} {employee.name}" if employee else "待對應", item["title"],
                         *(item[key] for key in PAY_FIELDS), item["gross"], item["deductions"], item["net"], item["status"],
                         "已對應" if employee else "待對應（未改員工資料）", item["note"], item["source"]])
        summary = {"薪資人數": len(rows), "Excel實領合計": money_sum(item["net"] for item in source_rows("payroll")),
                   "尚未對應": sum(not by_code.get(mapping.get(item["name"])) for item in source_rows("payroll"))}
    elif kind == "roster":
        days = list(range(1, last.day + 1))
        fields = ["姓名", "系統員工", "資料來源", *(f"{day}日" for day in days), "排休/休假", "特休", "病假", "事假", "例休/例假", "備註"]
        leaves = session.exec(select(LeaveRequest).where(LeaveRequest.start_date <= last, LeaveRequest.end_date >= first, LeaveRequest.status == LeaveStatus.approved)).all()
        start_utc = datetime.combine(first, datetime.min.time(), TZ).astimezone(timezone.utc).replace(tzinfo=None)
        end_utc = datetime.combine(last + timedelta(days=1), datetime.min.time(), TZ).astimezone(timezone.utc).replace(tzinfo=None)
        attendance = session.exec(select(AttendanceEvent).where(AttendanceEvent.happened_at >= start_utc, AttendanceEvent.happened_at < end_utc)).all()
        live = defaultdict(lambda: defaultdict(set))
        for item in attendance:
            live[item.employee_id][local_time(item.happened_at).day].add(item.event_type)
        for item in leaves:
            cursor = max(first, item.start_date)
            while cursor <= min(last, item.end_date):
                live[item.employee_id][cursor.day].add(f"核准{item.leave_type}")
                cursor += timedelta(days=1)
        shown = set()
        for item in source_rows("roster"):
            employee = by_code.get(mapping.get(item["name"]))
            counts = Counter(item["marks"].values())
            rows.append([item["name"], employee.employee_code if employee else "待對應", item["source"],
                         *(item["marks"].get(str(day), "") for day in days), counts["排休"] + counts["休假"], counts["特休"],
                         counts["病假"], counts["事假"], counts["例休"] + counts["例假"], item["note"]])
            if employee and employee.id not in shown:
                shown.add(employee.id)
                rows.append([employee.name, employee.employee_code, "系統實際打卡/核准請假", *("、".join(sorted(live[employee.id][day])) for day in days), None, None, None, None, None, "空白不代表休假或未出勤"])
        for employee in employees:
            if employee.id in live and employee.id not in shown:
                rows.append([employee.name, employee.employee_code, "系統實際打卡/核准請假", *("、".join(sorted(live[employee.id][day])) for day in days), None, None, None, None, None, "Excel尚無對應列"])
        summary = {"Excel排休人數": len(source_rows("roster")), "系統打卡事件": len(attendance), "核准請假筆數": len(leaves)}
    elif kind == "calendar":
        fields = ["日期", "時間", "類型", "標題/內容", "地點", "參與人員", "狀態", "備註", "來源"]
        rows = [[item["date"], item["time"], item["type"], item["title"], item["location"], item["people"], item["status"], item["note"], item["source"]] for item in source_rows("calendar")]
        sites = {item.id: item.name for item in session.exec(select(Worksite)).all()}
        member_codes = {item.id: item.employee_code for item in employees}
        assignments = session.exec(select(WorkAssignment).where(WorkAssignment.work_date >= first, WorkAssignment.work_date <= last)).all()
        members = session.exec(select(AssignmentMember).where(AssignmentMember.is_active.is_(True))).all()
        for item in assignments:
            rows.append([item.work_date.isoformat(), item.start_time.isoformat() if item.start_time else "", "系統派工", item.work_item,
                         sites.get(item.site_id), "、".join(member_codes.get(member.employee_id, "") for member in members if member.assignment_id == item.id),
                         item.status.value, item.notes, f"WorkAssignment #{item.id}"])
        meetings = session.exec(select(MeetingRecord).where(MeetingRecord.meeting_at >= datetime.combine(first, datetime.min.time()), MeetingRecord.meeting_at < datetime.combine(last + timedelta(days=1), datetime.min.time()))).all()
        for item in meetings:
            rows.append([item.meeting_at.date().isoformat(), item.meeting_at.strftime("%H:%M"), "系統會議", item.title, item.location, "、".join(item.attendee_codes), item.status.value, item.agenda, f"MeetingRecord #{item.id}"])
        leaves = session.exec(select(LeaveRequest).where(LeaveRequest.start_date <= last, LeaveRequest.end_date >= first, LeaveRequest.status == LeaveStatus.approved)).all()
        employee_ids = {item.id: item for item in employees}
        for item in leaves:
            cursor = max(first, item.start_date)
            while cursor <= min(last, item.end_date):
                rows.append([cursor.isoformat(), "", "核准請假", item.leave_type, "", employee_ids[item.employee_id].name if item.employee_id in employee_ids else "", "已核准", item.reason, f"LeaveRequest #{item.id}"])
                cursor += timedelta(days=1)
        rows.sort(key=lambda item: (item[0], item[1]))
        summary = {"Excel事項": len(source_rows("calendar")), "系統派工": len(assignments), "系統會議": len(meetings), "核准請假筆數": len(leaves)}
    elif kind == "annual":
        fields = ["月份", "系統收入", "系統支出", "系統淨額", "Excel原始收入（含結餘）", "Excel原始支出（含轉帳）", "Excel實領薪資", "薪資人數", "排休/休假", "特休", "病假", "事假", "例休/例假", "Excel行事曆事項"]
        for m in range(1, 13):
            period = f"{year}-{m:02d}"
            system = [item for item in finance if item.entry_date.month == m]
            source = [item for item in year_source_finance if item["month"] == period]
            salary = [item for item in data["payroll"] if item["month"] == period]
            marks = Counter(mark for item in data["roster"] if item["month"] == period for mark in item["marks"].values())
            income = money_sum(finance_values(item)[0] for item in system)
            expense = money_sum(finance_values(item)[1] for item in system)
            rows.append([period, income, expense, round(income - expense, 2), money_sum(item["income"] for item in source),
                         money_sum(item["expense"] for item in source), money_sum(item["net"] for item in salary) if salary else None,
                         len(salary), marks["排休"] + marks["休假"], marks["特休"], marks["病假"], marks["事假"],
                         marks["例休"] + marks["例假"], sum(item["month"] == period for item in data["calendar"])])
        summary = {"年度系統收入": money_sum(item[1] for item in rows), "年度系統支出": money_sum(item[2] for item in rows)}
    elif kind == "categories":
        fields = ["類別", "系統收入", "系統支出", "系統淨額", "Excel原始收入", "Excel原始支出", "Excel列數"]
        names = sorted({item.category for item in finance} | {item["category"] or item["status"] for item in year_source_finance})
        for name in names:
            system = [item for item in finance if item.category == name]
            source = [item for item in year_source_finance if (item["category"] or item["status"]) == name]
            income, expense = (money_sum(finance_values(item)[index] for item in system) for index in (0, 1))
            rows.append([name, income, expense, round(income - expense, 2), money_sum(item["income"] for item in source), money_sum(item["expense"] for item in source), len(source)])
    else:
        fields = ["公司/姓名", "統編/員工編號", "類別", "部門", "職稱", "聯絡人", "電話", "手機", "Email", "地址", "對應狀態", "來源"]
        contacts = session.exec(select(BusinessContact).order_by(BusinessContact.name)).all()
        keys = {item.normalized_key: item for item in contacts}
        tax_ids = {item.tax_id: item for item in contacts if item.tax_id}
        used = set()
        for item in data["contacts"]:
            target = tax_ids.get(item["tax_id"]) or keys.get(item["normalized_key"])
            if target:
                used.add(target.id)
            values = target.model_dump() if target else item
            rows.append([*(values.get(key) for key in ("name", "tax_id", "category", "department", "title", "contact_person", "phone", "mobile", "email", "address")), "已對應系統（以系統值為準）" if target else "Excel資料，尚未匯入通訊錄", item["source"]])
        for item in contacts:
            if item.id not in used:
                rows.append([item.name, item.tax_id, item.category, item.department, item.title, item.contact_person, item.phone, item.mobile, item.email, item.address, "系統已存", f"BusinessContact #{item.id}"])
        for item in employees:
            rows.append([item.name, item.employee_code, "系統員工", item.department, item.title, "", item.phone, item.phone, item.email, "", item.status.value, "Employee"])
        summary = {"系統公司通訊錄": len(contacts), "系統員工": len(employees), "Excel通訊錄": len(data["contacts"])}
    return {"kind": kind, "title": KINDS[kind], "month": month, "fields": fields, "rows": rows, "summary": summary,
            "snapshot_id": snapshot.id if snapshot else None,
            "notes": ["Excel版本與系統值分開保存；未入帳資料不計入系統收入支出。",
                      "薪資依原表F:K加給、L:P扣款計算，未依打卡自動扣薪或自動付款。",
                      "排休原註記完整保留；空白不等於休假，核准請假與實際打卡另列。"]}
